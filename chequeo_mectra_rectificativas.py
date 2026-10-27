from google.colab import drive
drive.mount('/content/drive')

import os
import re
import glob
import numpy as np
import pandas as pd
from datetime import date
from datetime import datetime

import openpyxl
from openpyxl import load_workbook

ruta_archivo_control = "/content/drive/My Drive/BonosLEC/Archivos procesamiento/Mectra/Control Bonos Emitidos.xlsx"

control = pd.read_excel(ruta_archivo_control, sheet_name="Hoja 1")
control.rename(columns={"Razón Social": "razon_social"}
, inplace=True)

fecha_generacion = datetime.now().strftime("%Y") + datetime.now().strftime("%m") + datetime.now().strftime("%d")

carpeta_mectra = "/content/drive/My Drive/BonosLEC/Archivos procesamiento/Mectra"
archivos_mectra = sorted(glob.glob(os.path.join(carpeta_mectra, "lps_*.csv")))

reporte_final = pd.DataFrame()
cuit_control = control['CUIT'].tolist()

for ruta_archivo_mectra in archivos_mectra:
  nombre_archivo = os.path.basename(ruta_archivo_mectra)
  match = re.search(r"lps_(\d{6})\.csv$", nombre_archivo)
  if not match:
    continue

  periodo_str = match.group(1)
  periodo = int(periodo_str)
  col_emision = int(periodo_str[2:])

  mectra = pd.read_csv(ruta_archivo_mectra, sep=',')
  mectra = mectra.drop(columns=['clae6', 'cuil', 'cod_bosocial', 'conss','conos', 'rem8', 'imp_ad_os', 'mes', 'remasigf'])

  reporte = pd.DataFrame([mectra.cuit]).transpose()
  reporte = reporte.drop_duplicates()
  reporte["index"] = [x for x in range(len(reporte))]
  reporte.set_index("index", inplace=True)
  reporte["razon_social"] = ""
  reporte["periodo"] = periodo

  reporte["total_emp"] = ""
  reporte["total_emp_lps_40"] = ""
  reporte["total_emp_lps_49"] = ""
  reporte["total_emp_lec_132"] = ""
  reporte["total_emp_lec_133"] = ""
  reporte["emision"] = ""

  for x in range(len(reporte)):
    cuit = reporte.loc[x, "cuit"]
    reporte.loc[x, "total_emp"] = len(mectra.loc[mectra["cuit"] == cuit])
    reporte.loc[x, "total_emp_lps_40"] = len(mectra.loc[(mectra["cuit"] == cuit) & (mectra["act_trab"] == 40)])
    reporte.loc[x, "total_emp_lps_49"] = len(mectra.loc[(mectra["cuit"] == cuit) & (mectra["act_trab"] == 49)])
    reporte.loc[x, "total_emp_lec_132"] = len(mectra.loc[(mectra["cuit"] == cuit) & (mectra["act_trab"] == 132)])
    reporte.loc[x, "total_emp_lec_133"] = len(mectra.loc[(mectra["cuit"] == cuit) & (mectra["act_trab"] == 133)])
    if cuit in cuit_control:
      reporte.loc[x, "razon_social"] = control.loc[control['CUIT'] == cuit,"razon_social"].values[0]
      reporte.loc[x, "emision"] = control.loc[control['CUIT'] == cuit][col_emision].values[0]
    if (reporte.loc[x, "total_emp_lec_132"] + reporte.loc[x, "total_emp_lec_133"]) == 0:
      reporte = reporte.drop(x)

  if reporte_final.empty:
    reporte_final = reporte
  else:
    reporte_final = pd.concat([reporte_final, reporte], ignore_index=True)

reporte_final.rename(columns={"cuit": "CUIT", "razon_social": "Razón Social", "total_emp": "Total Empleados", "total_emp_lps_40": "Cant Emp LPS 40", "total_emp_lps_49": "Cant Emp LPS 49",
                        "total_emp_lec_132": "Cant Emp LEC 132", "total_emp_lec_133": "Cant Emp LEC 133", "emision": "Emitido", "periodo": "Período"}
, inplace=True)
reporte_final.sort_values(by=["CUIT", "Período"], ascending=True, inplace=True)

nombre_reporte = "Reporte Archivos Mectra - " + fecha_generacion + ".xlsx"
ruta_archivo_salida = "/content/drive/My Drive/BonosLEC/Archivos Bonos LEC/Reporte/Mectra/" + nombre_reporte

writer = pd.ExcelWriter(ruta_archivo_salida, engine='xlsxwriter',
                        date_format='YYYY-MM-DD', datetime_format='YYYY-MM-DD')
df = reporte_final
df.to_excel(writer, sheet_name='Sheet1', startrow=1, header=False, index=False)

# Get the xlsxwriter workbook and worksheet objects.
workbook = writer.book
worksheet = writer.sheets['Sheet1']

# Get the dimensions of the dataframe.
(max_row, max_col) = df.shape

# Create a list of column headers, to use in add_table().
column_settings = [{'header': column} for column in df.columns]

# Add the Excel table structure. Pandas will add the data.
worksheet.add_table(0, 0, max_row, max_col - 1,
                    {'banded_rows': False, 'columns': column_settings})

# Make the columns wider for clarity.
worksheet.set_column(0, max_col - 1, 14)

# Close the Pandas Excel writer and output the Excel file.
writer.close()
